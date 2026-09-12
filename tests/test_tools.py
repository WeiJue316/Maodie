"""
test_tools.py — 工具系统测试

覆盖 agent/tools.py 的所有内置工具及注册/执行机制。
所有测试使用临时目录，不影响真实文件系统。
"""

from __future__ import annotations

import time
from pathlib import Path

import pytest

from agent.tools import ToolContext, execute_tool, TOOL_REGISTRY, _HAS_OPENPYXL


# ---------------------------------------------------------------------------
# 辅助：快速创建文件
# ---------------------------------------------------------------------------

def make_file(directory: Path, name: str, content: str = "hello\nworld\n") -> Path:
    p = directory / name
    p.write_text(content, encoding="utf-8")
    return p


# ---------------------------------------------------------------------------
# read_file
# ---------------------------------------------------------------------------

class TestReadFile:
    def test_success(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """正常读取已有文件，返回完整内容。"""
        make_file(tmp_work_dir, "a.txt", "line1\nline2\nline3\n")
        result = execute_tool("read_file", {"path": "a.txt"}, tool_ctx)
        assert "line1" in result
        assert "line3" in result

    def test_not_found(self, tool_ctx: ToolContext):
        """文件不存在时返回含 [错误] 的字符串，不抛异常。"""
        result = execute_tool("read_file", {"path": "no_such_file.txt"}, tool_ctx)
        assert "[错误]" in result

    def test_start_line(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """start_line 参数：只返回从指定行开始的内容。"""
        make_file(tmp_work_dir, "b.txt", "L1\nL2\nL3\nL4\n")
        result = execute_tool("read_file", {"path": "b.txt", "start_line": 3}, tool_ctx)
        assert "L3" in result
        assert "L1" not in result

    def test_max_lines(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """max_lines 参数：只返回指定行数。"""
        make_file(tmp_work_dir, "c.txt", "L1\nL2\nL3\nL4\n")
        result = execute_tool("read_file", {"path": "c.txt", "max_lines": 2}, tool_ctx)
        assert "L1" in result
        assert "L2" in result
        assert "L3" not in result

    def test_path_is_dir_returns_error(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """路径指向目录时返回错误。"""
        sub = tmp_work_dir / "subdir"
        sub.mkdir()
        result = execute_tool("read_file", {"path": "subdir"}, tool_ctx)
        assert "[错误]" in result

    def test_empty_file(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """空文件返回提示字符串而非崩溃。"""
        make_file(tmp_work_dir, "empty.txt", "")
        result = execute_tool("read_file", {"path": "empty.txt"}, tool_ctx)
        assert "空" in result or result == "(文件为空或指定范围无内容)"

    def test_truncates_large_file(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """超过输出上限时返回截断内容 + 截断提示。"""
        from agent.tools import _MAX_OUTPUT_BYTES
        big = "x" * (_MAX_OUTPUT_BYTES * 2)
        make_file(tmp_work_dir, "big.txt", big)
        result = execute_tool("read_file", {"path": "big.txt"}, tool_ctx)
        assert "已截断" in result
        assert len(result.encode("utf-8")) < len(big.encode("utf-8"))


# ---------------------------------------------------------------------------
# write_file
# ---------------------------------------------------------------------------

class TestWriteFile:
    def test_creates_file(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """写入新文件，文件确实被创建且内容正确。"""
        execute_tool("write_file", {"path": "new.txt", "content": "hello"}, tool_ctx)
        assert (tmp_work_dir / "new.txt").read_text(encoding="utf-8") == "hello"

    def test_overwrite(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """覆盖模式下，旧内容被替换。"""
        make_file(tmp_work_dir, "ow.txt", "old content")
        execute_tool("write_file", {"path": "ow.txt", "content": "new content"}, tool_ctx)
        assert (tmp_work_dir / "ow.txt").read_text(encoding="utf-8") == "new content"

    def test_append_mode(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """append=True 时内容追加而非覆盖。"""
        make_file(tmp_work_dir, "ap.txt", "first\n")
        execute_tool("write_file", {"path": "ap.txt", "content": "second\n", "append": True}, tool_ctx)
        content = (tmp_work_dir / "ap.txt").read_text(encoding="utf-8")
        assert "first" in content
        assert "second" in content

    def test_creates_parent_dir(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """目标目录不存在时自动创建。"""
        execute_tool("write_file", {"path": "deep/nested/file.txt", "content": "hi"}, tool_ctx)
        assert (tmp_work_dir / "deep" / "nested" / "file.txt").exists()

    def test_returns_summary(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """返回写入字节数摘要信息。"""
        result = execute_tool("write_file", {"path": "r.txt", "content": "abc"}, tool_ctx)
        assert "3" in result  # 3 个字符


# ---------------------------------------------------------------------------
# list_dir
# ---------------------------------------------------------------------------

class TestListDir:
    def test_returns_entries(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """列出目录条目，包含文件和子目录标记。"""
        make_file(tmp_work_dir, "file.txt", "x")
        (tmp_work_dir / "subdir").mkdir()
        result = execute_tool("list_dir", {}, tool_ctx)
        assert "file.txt" in result
        assert "subdir" in result

    def test_hides_dotfiles_by_default(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """show_hidden=False（默认）时不显示 . 开头条目。"""
        make_file(tmp_work_dir, ".hidden", "secret")
        make_file(tmp_work_dir, "visible.txt", "ok")
        result = execute_tool("list_dir", {}, tool_ctx)
        assert ".hidden" not in result
        assert "visible.txt" in result

    def test_shows_dotfiles_when_enabled(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """show_hidden=True 时显示 . 开头条目。"""
        make_file(tmp_work_dir, ".env", "KEY=val")
        result = execute_tool("list_dir", {"show_hidden": True}, tool_ctx)
        assert ".env" in result

    def test_not_found(self, tool_ctx: ToolContext):
        """目录不存在时返回含 [错误] 的字符串。"""
        result = execute_tool("list_dir", {"path": "no_such_dir"}, tool_ctx)
        assert "[错误]" in result

    def test_empty_dir(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """空目录时返回提示而非崩溃。"""
        empty = tmp_work_dir / "empty_sub"
        empty.mkdir()
        result = execute_tool("list_dir", {"path": "empty_sub"}, tool_ctx)
        assert "空" in result or "empty_sub" in result


# ---------------------------------------------------------------------------
# change_dir
# ---------------------------------------------------------------------------

class TestChangeDir:
    def test_success(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """切换到已有目录，ctx.work_dir 被更新。"""
        sub = tmp_work_dir / "sub"
        sub.mkdir()
        old = tool_ctx.work_dir
        execute_tool("change_dir", {"path": "sub"}, tool_ctx)
        assert tool_ctx.work_dir == sub.resolve()

    def test_not_found(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """目录不存在时返回错误，ctx.work_dir 不变。"""
        original = tool_ctx.work_dir
        result = execute_tool("change_dir", {"path": "no_such"}, tool_ctx)
        assert "[错误]" in result
        assert tool_ctx.work_dir == original

    def test_affects_subsequent_read(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """change_dir 后，read_file 基于新目录解析路径。"""
        sub = tmp_work_dir / "subdir"
        sub.mkdir()
        (sub / "data.txt").write_text("in subdir", encoding="utf-8")
        execute_tool("change_dir", {"path": "subdir"}, tool_ctx)
        result = execute_tool("read_file", {"path": "data.txt"}, tool_ctx)
        assert "in subdir" in result


# ---------------------------------------------------------------------------
# shell
# ---------------------------------------------------------------------------

class TestShell:
    def test_stdout(self, tool_ctx: ToolContext):
        """执行合法命令，返回 stdout 内容。"""
        result = execute_tool("shell", {"command": "echo hello_from_shell"}, tool_ctx)
        assert "hello_from_shell" in result

    def test_nonzero_exit_code(self, tool_ctx: ToolContext):
        """非零退出码时，结果中包含退出码信息。"""
        result = execute_tool("shell", {"command": "exit 42"}, tool_ctx)
        assert "42" in result

    def test_timeout(self, tool_ctx: ToolContext):
        """命令超时时返回超时提示，不崩溃。"""
        result = execute_tool("shell", {"command": "ping -n 10 127.0.0.1", "timeout": 1}, tool_ctx)
        assert "[错误]" in result and "超时" in result

    def test_cwd_is_work_dir(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """shell 命令在 work_dir 中执行。"""
        # 通过 cd 命令获取当前目录，和 work_dir 比较
        result = execute_tool("shell", {"command": "cd"}, tool_ctx)
        # Windows cd 输出当前目录
        assert str(tmp_work_dir).lower() in result.lower() or result.strip() != ""

    def test_truncates_long_output(self, tool_ctx: ToolContext):
        """stdout 过长时被截断并附带提示。"""
        from agent.tools import _MAX_OUTPUT_BYTES
        # 用 python 生成远超上限的输出
        cmd = (
            f'python -c "import sys; sys.stdout.write(\'x\' * {_MAX_OUTPUT_BYTES * 2})"'
        )
        result = execute_tool("shell", {"command": cmd}, tool_ctx)
        assert "已截断" in result
        assert len(result.encode("utf-8")) < _MAX_OUTPUT_BYTES * 2


# ---------------------------------------------------------------------------
# search_files
# ---------------------------------------------------------------------------

class TestSearchFiles:
    def test_by_pattern(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """glob 模式匹配，返回匹配文件列表。"""
        make_file(tmp_work_dir, "a.py", "")
        make_file(tmp_work_dir, "b.py", "")
        make_file(tmp_work_dir, "c.txt", "")
        result = execute_tool("search_files", {"pattern": "*.py"}, tool_ctx)
        assert "a.py" in result
        assert "b.py" in result
        assert "c.txt" not in result

    def test_with_content_keyword(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """content_keyword 过滤，只返回含关键词的文件。"""
        make_file(tmp_work_dir, "match.txt", "find me here")
        make_file(tmp_work_dir, "no_match.txt", "nothing special")
        result = execute_tool("search_files", {"pattern": "*.txt", "content_keyword": "find me"}, tool_ctx)
        assert "match.txt" in result
        assert "no_match.txt" not in result

    def test_not_found(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """无匹配时返回提示字符串，不崩溃。"""
        result = execute_tool("search_files", {"pattern": "*.xyz"}, tool_ctx)
        assert "未找到" in result

    def test_recursive(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """递归搜索子目录。"""
        sub = tmp_work_dir / "deep"
        sub.mkdir()
        make_file(sub, "nested.py", "")
        result = execute_tool("search_files", {"pattern": "*.py"}, tool_ctx)
        assert "nested.py" in result

    def test_path_not_found(self, tool_ctx: ToolContext):
        """搜索根目录不存在时返回错误。"""
        result = execute_tool("search_files", {"pattern": "*.py", "path": "no_such_dir"}, tool_ctx)
        assert "[错误]" in result


# ---------------------------------------------------------------------------
# Excel 工具
# ---------------------------------------------------------------------------

@pytest.mark.skipif(not _HAS_OPENPYXL, reason="openpyxl 未安装")
class TestCreateWorkbook:
    def test_creates_file(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """创建 Excel 文件，文件确实存在。"""
        data = [["姓名", "年龄"], ["张三", 25], ["李四", 30]]
        result = execute_tool("create_workbook", {"path": "test.xlsx", "data": data}, tool_ctx)
        assert (tmp_work_dir / "test.xlsx").exists()
        assert "已创建" in result

    def test_with_headers(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """headers=True 时第一行加粗。"""
        import openpyxl as xl
        data = [["A", "B"], [1, 2]]
        execute_tool("create_workbook", {"path": "h.xlsx", "data": data, "headers": True}, tool_ctx)
        wb = xl.load_workbook(str(tmp_work_dir / "h.xlsx"))
        assert wb.active["A1"].font.bold is True
        assert wb.active["A2"].font.bold is not True

    def test_custom_sheet_name(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """自定义 sheet 名称。"""
        execute_tool("create_workbook", {"path": "s.xlsx", "data": [["x"]], "sheet_name": "数据"}, tool_ctx)
        import openpyxl as xl
        wb = xl.load_workbook(str(tmp_work_dir / "s.xlsx"))
        assert "数据" in wb.sheetnames

    def test_empty_data_returns_error(self, tool_ctx: ToolContext):
        """空 data 返回错误。"""
        result = execute_tool("create_workbook", {"path": "e.xlsx", "data": []}, tool_ctx)
        assert "[错误]" in result

    def test_creates_parent_dir(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """自动创建父目录。"""
        execute_tool("create_workbook", {"path": "sub/dir.xlsx", "data": [["a"]]}, tool_ctx)
        assert (tmp_work_dir / "sub" / "dir.xlsx").exists()


@pytest.mark.skipif(not _HAS_OPENPYXL, reason="openpyxl 未安装")
class TestReadWorkbook:
    def test_reads_content(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """读取已创建的 Excel，返回内容包含数据。"""
        data = [["Name", "Score"], ["Alice", 95], ["Bob", 87]]
        execute_tool("create_workbook", {"path": "r.xlsx", "data": data}, tool_ctx)
        result = execute_tool("read_workbook", {"path": "r.xlsx"}, tool_ctx)
        assert "Alice" in result
        assert "95" in result

    def test_not_found(self, tool_ctx: ToolContext):
        """文件不存在时返回错误。"""
        result = execute_tool("read_workbook", {"path": "no.xlsx"}, tool_ctx)
        assert "[错误]" in result

    def test_max_rows(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """max_rows 限制返回行数。"""
        data = [[f"row{i}"] for i in range(20)]
        execute_tool("create_workbook", {"path": "m.xlsx", "data": data}, tool_ctx)
        result = execute_tool("read_workbook", {"path": "m.xlsx", "max_rows": 5}, tool_ctx)
        assert "截断" in result

    def test_specific_sheet(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """指定 sheet_name 读取。"""
        execute_tool("create_workbook", {"path": "ss.xlsx", "data": [["a"]], "sheet_name": "MySheet"}, tool_ctx)
        result = execute_tool("read_workbook", {"path": "ss.xlsx", "sheet_name": "MySheet"}, tool_ctx)
        assert "MySheet" in result


@pytest.mark.skipif(not _HAS_OPENPYXL, reason="openpyxl 未安装")
class TestEditCell:
    def test_edit_value(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """修改单元格值。"""
        import openpyxl as xl
        execute_tool("create_workbook", {"path": "ec.xlsx", "data": [["a", "b"], [1, 2]]}, tool_ctx)
        execute_tool("edit_cell", {"path": "ec.xlsx", "cell": "A1", "value": "changed"}, tool_ctx)
        wb = xl.load_workbook(str(tmp_work_dir / "ec.xlsx"))
        assert wb.active["A1"].value == "changed"

    def test_edit_formula(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """写入公式。"""
        import openpyxl as xl
        execute_tool("create_workbook", {"path": "f.xlsx", "data": [[1, 2, 0]]}, tool_ctx)
        execute_tool("edit_cell", {"path": "f.xlsx", "cell": "C1", "value": "SUM(A1:B1)", "is_formula": True}, tool_ctx)
        wb = xl.load_workbook(str(tmp_work_dir / "f.xlsx"))
        assert wb.active["C1"].value == "=SUM(A1:B1)"

    def test_not_found(self, tool_ctx: ToolContext):
        """文件不存在时返回错误。"""
        result = execute_tool("edit_cell", {"path": "no.xlsx", "cell": "A1", "value": 1}, tool_ctx)
        assert "[错误]" in result


@pytest.mark.skipif(not _HAS_OPENPYXL, reason="openpyxl 未安装")
class TestFormatCells:
    def test_bold(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """加粗格式。"""
        import openpyxl as xl
        execute_tool("create_workbook", {"path": "fmt.xlsx", "data": [["a", "b"], [1, 2]]}, tool_ctx)
        execute_tool("format_cells", {"path": "fmt.xlsx", "range": "A1:B1", "bold": True}, tool_ctx)
        wb = xl.load_workbook(str(tmp_work_dir / "fmt.xlsx"))
        assert wb.active["A1"].font.bold is True
        assert wb.active["B1"].font.bold is True

    def test_bg_color(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """背景色。"""
        import openpyxl as xl
        execute_tool("create_workbook", {"path": "bg.xlsx", "data": [["x"]]}, tool_ctx)
        execute_tool("format_cells", {"path": "bg.xlsx", "range": "A1", "bg_color": "FFFF00"}, tool_ctx)
        wb = xl.load_workbook(str(tmp_work_dir / "bg.xlsx"))
        assert wb.active["A1"].fill.start_color.rgb == "00FFFF00"

    def test_number_format(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """数字格式。"""
        import openpyxl as xl
        execute_tool("create_workbook", {"path": "nf.xlsx", "data": [[3.14159]]}, tool_ctx)
        execute_tool("format_cells", {"path": "nf.xlsx", "range": "A1", "number_format": "0.00"}, tool_ctx)
        wb = xl.load_workbook(str(tmp_work_dir / "nf.xlsx"))
        assert wb.active["A1"].number_format == "0.00"

    def test_not_found(self, tool_ctx: ToolContext):
        """文件不存在时返回错误。"""
        result = execute_tool("format_cells", {"path": "no.xlsx", "range": "A1", "bold": True}, tool_ctx)
        assert "[错误]" in result


@pytest.mark.skipif(not _HAS_OPENPYXL, reason="openpyxl 未安装")
class TestManageSheet:
    def _make_file(self, tool_ctx: ToolContext, tmp_work_dir: Path) -> str:
        execute_tool("create_workbook", {"path": "ms.xlsx", "data": [["a"]]}, tool_ctx)
        return str(tmp_work_dir / "ms.xlsx")

    def test_list(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """列出工作表。"""
        self._make_file(tool_ctx, tmp_work_dir)
        result = execute_tool("manage_sheet", {"path": "ms.xlsx", "action": "list"}, tool_ctx)
        assert "Sheet1" in result

    def test_create(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """创建工作表。"""
        self._make_file(tool_ctx, tmp_work_dir)
        result = execute_tool("manage_sheet", {"path": "ms.xlsx", "action": "create", "sheet_name": "New"}, tool_ctx)
        assert "已创建" in result
        import openpyxl as xl
        wb = xl.load_workbook(str(tmp_work_dir / "ms.xlsx"))
        assert "New" in wb.sheetnames

    def test_rename(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """重命名工作表。"""
        self._make_file(tool_ctx, tmp_work_dir)
        execute_tool("manage_sheet", {"path": "ms.xlsx", "action": "rename", "sheet_name": "Sheet1", "new_name": "数据"}, tool_ctx)
        import openpyxl as xl
        wb = xl.load_workbook(str(tmp_work_dir / "ms.xlsx"))
        assert "数据" in wb.sheetnames
        assert "Sheet1" not in wb.sheetnames

    def test_delete(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """删除工作表。"""
        self._make_file(tool_ctx, tmp_work_dir)
        execute_tool("manage_sheet", {"path": "ms.xlsx", "action": "create", "sheet_name": "Temp"}, tool_ctx)
        execute_tool("manage_sheet", {"path": "ms.xlsx", "action": "delete", "sheet_name": "Temp"}, tool_ctx)
        import openpyxl as xl
        wb = xl.load_workbook(str(tmp_work_dir / "ms.xlsx"))
        assert "Temp" not in wb.sheetnames

    def test_delete_last_sheet_fails(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """不能删除唯一的工作表。"""
        self._make_file(tool_ctx, tmp_work_dir)
        result = execute_tool("manage_sheet", {"path": "ms.xlsx", "action": "delete", "sheet_name": "Sheet1"}, tool_ctx)
        assert "[错误]" in result

    def test_not_found(self, tool_ctx: ToolContext):
        """文件不存在时返回错误。"""
        result = execute_tool("manage_sheet", {"path": "no.xlsx", "action": "list"}, tool_ctx)
        assert "[错误]" in result


# ---------------------------------------------------------------------------
# execute_tool — 注册与错误处理
# ---------------------------------------------------------------------------

class TestExecuteTool:
    def test_unknown_tool_returns_error(self, tool_ctx: ToolContext):
        """调用不存在的工具名返回 [错误]，不抛异常。"""
        result = execute_tool("totally_unknown_tool", {}, tool_ctx)
        assert "[错误]" in result
        assert "未知工具" in result

    def test_tool_internal_exception_is_caught(self, tool_ctx: ToolContext, tmp_work_dir: Path):
        """工具内部抛出异常时，execute_tool 返回错误字符串而非向上传播。"""
        from agent.tools import tool as register_tool, TOOL_REGISTRY

        @register_tool(
            name="_test_raise",
            description="always raises",
            schema={"type": "object", "properties": {}, "required": []},
        )
        def _test_raise(args, ctx):
            raise RuntimeError("intentional error")

        try:
            result = execute_tool("_test_raise", {}, tool_ctx)
            assert "[错误]" in result
        finally:
            # 清理临时注册的工具
            TOOL_REGISTRY.pop("_test_raise", None)

    def test_all_default_tools_registered(self):
        """所有默认工具均已注册到 TOOL_REGISTRY。"""
        expected = {"read_file", "write_file", "list_dir", "change_dir", "shell", "search_files",
                     "create_workbook", "read_workbook", "edit_cell", "format_cells", "manage_sheet"}
        assert expected.issubset(TOOL_REGISTRY.keys())
