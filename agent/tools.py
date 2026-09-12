"""
工具系统：注册机制 + 内置工具。

使用 @tool 装饰器注册工具，工具元信息（name/description/schema）内聚在函数上。
ToolContext 在执行时传入，工具可通过它读写运行时状态（当前工作目录等）。
"""

from __future__ import annotations

import fnmatch
import json
import os
import subprocess
import urllib.error
import urllib.request
from dataclasses import dataclass
from pathlib import Path
from typing import TYPE_CHECKING, Any, Callable

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    _HAS_OPENPYXL = True
except ImportError:
    _HAS_OPENPYXL = False

if TYPE_CHECKING:
    from .mcp import MCPManager
    from .memory import MemoryManager
    from .memory_search import MemorySearch
    from .orchestrator import Orchestrator
    from .skills import SkillManager


# ---------------------------------------------------------------------------
# ToolContext
# ---------------------------------------------------------------------------

@dataclass
class ToolContext:
    """运行时状态，在工具执行时传入。"""
    work_dir: Path          # 当前工作目录（可被 change_dir 修改）
    session_id: str = ""
    skill_manager: SkillManager | None = None
    memory_manager: MemoryManager | None = None
    memory_search: MemorySearch | None = None
    mcp_manager: MCPManager | None = None
    orchestrator: Orchestrator | None = None
    can_spawn_agents: bool = True  # 主 agent 为 True，子 agent 必须为 False

    def resolve(self, path: str) -> Path:
        """将相对路径解析为绝对路径（相对于 work_dir）。"""
        p = Path(path)
        if p.is_absolute():
            return p
        return (self.work_dir / p).resolve()


# ---------------------------------------------------------------------------
# 工具注册
# ---------------------------------------------------------------------------

@dataclass
class ToolDef:
    name: str
    description: str
    schema: dict[str, Any]          # OpenAI function parameters schema
    func: Callable[[dict, ToolContext], str]

    def to_openai_schema(self) -> dict[str, Any]:
        return {
            "type": "function",
            "function": {
                "name": self.name,
                "description": self.description,
                "parameters": self.schema,
            },
        }


TOOL_REGISTRY: dict[str, ToolDef] = {}


# 工具输出上限：防止超长文件/命令输出灌爆 LLM context
_MAX_OUTPUT_BYTES = 20_000


def _truncate_output(text: str, limit: int = _MAX_OUTPUT_BYTES) -> str:
    """按 UTF-8 字节数截断，超限时附加截断提示。"""
    encoded = text.encode("utf-8", errors="replace")
    if len(encoded) <= limit:
        return text
    truncated = encoded[:limit].decode("utf-8", errors="ignore")
    omitted = len(encoded) - len(truncated.encode("utf-8"))
    return (
        f"{truncated}\n"
        f"...[已截断，省略 {omitted} 字节；"
        f"如需完整内容请缩小范围，例如指定 start_line/max_lines]"
    )


def tool(
    name: str,
    description: str,
    schema: dict[str, Any],
) -> Callable:
    """装饰器：将函数注册为工具。函数签名必须为 (args: dict, ctx: ToolContext) -> str。"""
    def decorator(func: Callable) -> Callable:
        TOOL_REGISTRY[name] = ToolDef(
            name=name,
            description=description,
            schema=schema,
            func=func,
        )
        return func
    return decorator


def get_enabled_tools(enabled_names: list[str]) -> list[ToolDef]:
    """返回已启用的工具列表。"""
    return [TOOL_REGISTRY[n] for n in enabled_names if n in TOOL_REGISTRY]


def get_openai_schemas(enabled_names: list[str]) -> list[dict[str, Any]]:
    """返回 OpenAI function calling 格式的工具 schema 列表。

    包含两部分：
    1. enabled_names 中列出的内置工具
    2. TOOL_REGISTRY 中名称含 '__' 的工具（MCP 工具，自动包含）
    """
    schemas: list[dict[str, Any]] = []
    seen: set[str] = set()

    # 内置工具
    for t in get_enabled_tools(enabled_names):
        schemas.append(t.to_openai_schema())
        seen.add(t.name)

    # MCP 工具（名称含 __）
    for name, td in TOOL_REGISTRY.items():
        if name not in seen and "__" in name:
            schemas.append(td.to_openai_schema())

    return schemas


def execute_tool(name: str, arguments: dict[str, Any], ctx: ToolContext) -> str:
    """
    执行工具。

    Returns:
        工具执行结果字符串；出错时返回错误描述（不抛出异常）。
    """
    tool_def = TOOL_REGISTRY.get(name)
    if tool_def is None:
        return f"[错误] 未知工具：{name}"
    try:
        return tool_def.func(arguments, ctx)
    except Exception as e:
        return f"[错误] 工具 {name} 执行失败：{e}"


# ---------------------------------------------------------------------------
# 内置工具实现
# ---------------------------------------------------------------------------

@tool(
    name="read_file",
    description="读取文件内容并返回。支持指定起始行和最大行数。",
    schema={
        "type": "object",
        "properties": {
            "path": {
                "type": "string",
                "description": "文件路径（相对于工作目录或绝对路径）",
            },
            "encoding": {
                "type": "string",
                "description": "文件编码，默认 utf-8",
                "default": "utf-8",
            },
            "start_line": {
                "type": "integer",
                "description": "从第几行开始读取（从 1 计），默认 1",
                "default": 1,
            },
            "max_lines": {
                "type": "integer",
                "description": "最多读取多少行，默认不限制",
            },
        },
        "required": ["path"],
    },
)
def read_file(args: dict, ctx: ToolContext) -> str:
    path = ctx.resolve(args["path"])
    encoding = args.get("encoding", "utf-8")
    start_line = max(1, int(args.get("start_line", 1)))
    max_lines = args.get("max_lines")

    if not path.exists():
        return f"[错误] 文件不存在：{path}"
    if not path.is_file():
        return f"[错误] 路径不是文件：{path}"

    with open(path, encoding=encoding, errors="replace") as f:
        lines = f.readlines()

    selected = lines[start_line - 1:]
    if max_lines is not None:
        selected = selected[:int(max_lines)]

    if not selected:
        return "(文件为空或指定范围无内容)"
    return _truncate_output("".join(selected))


@tool(
    name="write_file",
    description="将内容写入文件。默认覆盖，可选追加模式。",
    schema={
        "type": "object",
        "properties": {
            "path": {
                "type": "string",
                "description": "文件路径（相对于工作目录或绝对路径）",
            },
            "content": {
                "type": "string",
                "description": "要写入的内容",
            },
            "append": {
                "type": "boolean",
                "description": "是否追加模式，默认 false（覆盖）",
                "default": False,
            },
            "encoding": {
                "type": "string",
                "description": "文件编码，默认 utf-8",
                "default": "utf-8",
            },
        },
        "required": ["path", "content"],
    },
)
def write_file(args: dict, ctx: ToolContext) -> str:
    path = ctx.resolve(args["path"])
    content = args["content"]
    append = args.get("append", False)
    encoding = args.get("encoding", "utf-8")

    path.parent.mkdir(parents=True, exist_ok=True)
    mode = "a" if append else "w"
    with open(path, mode, encoding=encoding) as f:
        f.write(content)

    action = "追加写入" if append else "写入"
    return f"已{action} {len(content)} 个字符到 {path}"


@tool(
    name="list_dir",
    description="列出目录下的文件和子目录。",
    schema={
        "type": "object",
        "properties": {
            "path": {
                "type": "string",
                "description": "目录路径，默认为当前工作目录",
                "default": ".",
            },
            "show_hidden": {
                "type": "boolean",
                "description": "是否显示以 . 开头的隐藏文件，默认 false",
                "default": False,
            },
        },
        "required": [],
    },
)
def list_dir(args: dict, ctx: ToolContext) -> str:
    path_str = args.get("path", ".")
    path = ctx.resolve(path_str)
    show_hidden = args.get("show_hidden", False)

    if not path.exists():
        return f"[错误] 路径不存在：{path}"
    if not path.is_dir():
        return f"[错误] 路径不是目录：{path}"

    entries = sorted(path.iterdir(), key=lambda e: (e.is_file(), e.name.lower()))
    lines = []
    for entry in entries:
        if not show_hidden and entry.name.startswith("."):
            continue
        if entry.is_dir():
            lines.append(f"[目录] {entry.name}/")
        else:
            size = entry.stat().st_size
            lines.append(f"[文件] {entry.name}  ({_human_size(size)})")

    if not lines:
        return f"目录为空：{path}"
    return f"目录：{path}\n" + "\n".join(lines)


@tool(
    name="change_dir",
    description="切换当前工作目录。后续工具调用的相对路径将基于新目录。",
    schema={
        "type": "object",
        "properties": {
            "path": {
                "type": "string",
                "description": "目标目录路径（相对或绝对）",
            },
        },
        "required": ["path"],
    },
)
def change_dir(args: dict, ctx: ToolContext) -> str:
    new_dir = ctx.resolve(args["path"])
    if not new_dir.exists():
        return f"[错误] 目录不存在：{new_dir}"
    if not new_dir.is_dir():
        return f"[错误] 路径不是目录：{new_dir}"
    old_dir = ctx.work_dir
    ctx.work_dir = new_dir
    return f"工作目录已从 {old_dir} 切换到 {new_dir}"


@tool(
    name="shell",
    description="在当前工作目录执行 shell 命令，返回标准输出和标准错误。",
    schema={
        "type": "object",
        "properties": {
            "command": {
                "type": "string",
                "description": "要执行的 shell 命令",
            },
            "timeout": {
                "type": "integer",
                "description": "超时秒数，默认 30",
                "default": 30,
            },
        },
        "required": ["command"],
    },
)
def shell(args: dict, ctx: ToolContext) -> str:
    command = args["command"]
    timeout = int(args.get("timeout", 30))

    try:
        result = subprocess.run(
            command,
            shell=True,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=timeout,
            cwd=str(ctx.work_dir),
        )
        output_parts = []
        if result.stdout:
            output_parts.append(result.stdout)
        if result.stderr:
            output_parts.append(f"[stderr]\n{result.stderr}")
        if result.returncode != 0:
            output_parts.append(f"[退出码] {result.returncode}")
        output = "\n".join(output_parts) if output_parts else "(命令执行成功，无输出)"
        return _truncate_output(output)
    except subprocess.TimeoutExpired:
        return f"[错误] 命令超时（{timeout}s）：{command}"


@tool(
    name="search_files",
    description="在目录中递归搜索匹配文件名 glob 模式的文件，可选在文件内容中搜索关键词。",
    schema={
        "type": "object",
        "properties": {
            "pattern": {
                "type": "string",
                "description": "文件名 glob 模式，如 '*.py' 或 'config*'",
            },
            "path": {
                "type": "string",
                "description": "搜索根目录，默认为当前工作目录",
                "default": ".",
            },
            "content_keyword": {
                "type": "string",
                "description": "可选：在文件内容中搜索的关键词",
            },
            "max_results": {
                "type": "integer",
                "description": "最多返回结果数，默认 50",
                "default": 50,
            },
        },
        "required": ["pattern"],
    },
)
def search_files(args: dict, ctx: ToolContext) -> str:
    pattern = args["pattern"]
    root = ctx.resolve(args.get("path", "."))
    keyword = args.get("content_keyword")
    max_results = int(args.get("max_results", 50))

    if not root.exists():
        return f"[错误] 路径不存在：{root}"

    matched: list[str] = []
    for dirpath, _dirs, filenames in os.walk(root):
        for filename in filenames:
            if fnmatch.fnmatch(filename, pattern):
                full_path = Path(dirpath) / filename
                if keyword:
                    try:
                        text = full_path.read_text(encoding="utf-8", errors="replace")
                        if keyword not in text:
                            continue
                    except Exception:
                        continue
                matched.append(str(full_path))
                if len(matched) >= max_results:
                    break
        if len(matched) >= max_results:
            break

    if not matched:
        desc = f"模式 '{pattern}'"
        if keyword:
            desc += f" 且包含 '{keyword}'"
        return f"未找到匹配文件（{desc}）"

    header = f"找到 {len(matched)} 个文件（搜索根目录：{root}）"
    if len(matched) >= max_results:
        header += f"（已达上限 {max_results}，可能有更多）"
    return header + "\n" + "\n".join(matched)


# ---------------------------------------------------------------------------
# Skill 工具
# ---------------------------------------------------------------------------


@tool(
    name="select_skill",
    description="获取指定 Skill 的文件路径和工具信息。使用 read_file 工具加载返回的路径以获取完整指令。",
    schema={
        "type": "object",
        "properties": {
            "name": {
                "type": "string",
                "description": "要查找的 Skill 名称",
            },
        },
        "required": ["name"],
    },
)
def select_skill(args: dict, ctx: ToolContext) -> str:
    name = args["name"]
    sm = ctx.skill_manager
    if sm is None:
        return "[错误] Skill 系统未初始化"

    skill = sm.get_skill(name)
    if skill is None:
        available = [s.name for s in sm.skills]
        return f"[错误] 未找到 Skill：{name}。可用 Skills：{', '.join(available)}"

    parts = [f"[Skill: {skill.name}]"]
    parts.append(f"文件路径: {skill.file_path}")

    if skill.tools:
        parts.append(f"可用工具: {', '.join(skill.tools)}")
    if skill.unavailable_tools:
        parts.append(f"⚠️ 以下工具不可用，请勿调用: {', '.join(skill.unavailable_tools)}")

    parts.append("\n请使用 read_file 工具加载上述文件获取详细指令。")
    return "\n".join(parts)


@tool(
    name="install_skill",
    description="从 URL 安装一个 Skill。支持直接链接和会自动跟随重定向。URL 以 / 结尾时会自动尝试 SKILL.md。",
    schema={
        "type": "object",
        "properties": {
            "url": {
                "type": "string",
                "description": "Skill 的 SKILL.md 文件 URL",
            },
            "name": {
                "type": "string",
                "description": "指定 Skill 目录名（可选，默认从 frontmatter 或 URL 推断）",
            },
        },
        "required": ["url"],
    },
)
def install_skill(args: dict, ctx: ToolContext) -> str:
    from .skills import _parse_frontmatter

    sm = ctx.skill_manager
    if sm is None:
        return "[错误] Skill 系统未初始化"

    url = args["url"].strip()
    override_name = args.get("name", "").strip()

    # URL 以 / 结尾 → 尝试 SKILL.md
    if url.endswith("/"):
        url = url + "SKILL.md"

    # 下载内容
    try:
        req = urllib.request.Request(url, headers={"User-Agent": "agent-skill-installer/1.0"})
        with urllib.request.urlopen(req, timeout=30) as resp:
            raw = resp.read()
            content = raw.decode("utf-8")
    except urllib.error.HTTPError as e:
        return f"[错误] HTTP {e.code}：{e.reason}（{url}）"
    except urllib.error.URLError as e:
        return f"[错误] 网络错误：{e.reason}（{url}）"
    except Exception as e:
        return f"[错误] 下载失败：{e}"

    # 验证 frontmatter
    meta, _ = _parse_frontmatter(content)
    if not meta.get("name") and not override_name:
        return "[错误] 内容缺少 frontmatter（---name: ...---），不是有效的 Skill 文件。可用 name 参数指定名称。"

    skill_name = override_name or meta["name"]

    # 创建目录并写入
    skills_dir = sm._dir
    skill_dir = skills_dir / skill_name
    skill_dir.mkdir(parents=True, exist_ok=True)
    skill_file = skill_dir / "SKILL.md"
    skill_file.write_text(content, encoding="utf-8")

    # 重新发现 skills
    sm.discover(available_tools=set(TOOL_REGISTRY.keys()))

    desc = meta.get("description", "(无描述)")
    action = "更新" if skill_file.exists() else "安装"
    return f"[成功] {action} Skill：{skill_name}\n描述：{desc}\n路径：{skill_file}"


# ---------------------------------------------------------------------------
# 记忆工具
# ---------------------------------------------------------------------------


@tool(
    name="write_memory",
    description="将信息写入长期记忆。记忆会跨会话保存，下次启动时自动加载。",
    schema={
        "type": "object",
        "properties": {
            "category": {
                "type": "string",
                "description": "记忆分类，如：用户偏好、项目约束、历史上下文",
            },
            "content": {
                "type": "string",
                "description": "要记住的具体内容，简洁明了",
            },
        },
        "required": ["category", "content"],
    },
)
def write_memory(args: dict, ctx: ToolContext) -> str:
    mm = ctx.memory_manager
    if mm is None:
        return "[错误] 长期记忆系统未初始化"
    category = args["category"]
    content = args["content"]
    mm.append(category, content)
    return f"已记住：[{category}] {content}"


@tool(
    name="search_memory",
    description="搜索历史记忆（observation），返回与查询相关的记忆条目。",
    schema={
        "type": "object",
        "properties": {
            "query": {
                "type": "string",
                "description": "搜索关键词或语义描述",
            },
            "limit": {
                "type": "integer",
                "description": "返回结果数量，默认 5",
                "default": 5,
            },
        },
        "required": ["query"],
    },
)
def search_memory(args: dict, ctx: ToolContext) -> str:
    ms = ctx.memory_search
    if ms is None:
        return "[错误] 记忆检索系统未初始化"
    query = args["query"]
    limit = int(args.get("limit", 5))
    observations = ms.search(query, limit=limit)
    if not observations:
        return "未找到相关记忆。"
    parts = [f"找到 {len(observations)} 条相关记忆：\n"]
    for obs in observations:
        parts.append(f"[{obs.type}] {obs.title}")
        parts.append(f"  {obs.narrative[:100]}")
        if obs.concepts:
            parts.append(f"  概念：{', '.join(obs.concepts)}")
        parts.append("")
    return "\n".join(parts)


@tool(
    name="get_daily_memory",
    description="查看指定日期的记忆（observation）。不传日期则查看今天。返回当天提取的所有记忆条目。",
    schema={
        "type": "object",
        "properties": {
            "date": {
                "type": "string",
                "description": "日期，格式 YYYY-MM-DD，默认今天",
            },
        },
        "required": [],
    },
)
def get_daily_memory(args: dict, ctx: ToolContext) -> str:
    ms = ctx.memory_search
    if ms is None:
        return "[错误] 记忆检索系统未初始化"
    from datetime import date as _date
    date_str = args.get("date") or _date.today().isoformat()
    from .observation import generate_daily_memory
    return generate_daily_memory(date_str, ms.store)


# ---------------------------------------------------------------------------
# Excel 工具
# ---------------------------------------------------------------------------

_EXCEL_ERROR = "[错误] 需要安装 openpyxl：pip install openpyxl"


@tool(
    name="create_workbook",
    description="创建一个新的 Excel (.xlsx) 文件并写入数据。支持表头加粗和自动列宽。",
    schema={
        "type": "object",
        "properties": {
            "path": {
                "type": "string",
                "description": "文件路径（相对于工作目录或绝对路径）",
            },
            "data": {
                "type": "array",
                "items": {"type": "array"},
                "description": "二维数组，每行是一个数组。如 [[\"姓名\",\"年龄\"],[\"张三\",25]]",
            },
            "sheet_name": {
                "type": "string",
                "description": "工作表名称，默认 Sheet1",
                "default": "Sheet1",
            },
            "headers": {
                "type": "boolean",
                "description": "第一行是否作为表头（加粗），默认 false",
                "default": False,
            },
        },
        "required": ["path", "data"],
    },
)
def create_workbook(args: dict, ctx: ToolContext) -> str:
    if not _HAS_OPENPYXL:
        return _EXCEL_ERROR

    path = ctx.resolve(args["path"])
    data = args["data"]
    sheet_name = args.get("sheet_name", "Sheet1")
    headers = args.get("headers", False)

    if not data:
        return "[错误] data 不能为空"

    path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    bold_font = Font(bold=True) if headers else None

    for row_idx, row in enumerate(data, start=1):
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if headers and row_idx == 1 and bold_font:
                cell.font = bold_font

    # 自动列宽
    for col_idx in range(1, len(data[0]) + 1 if data else 1):
        max_len = 0
        for row in data:
            if col_idx <= len(row):
                cell_len = len(str(row[col_idx - 1]))
                max_len = max(max_len, cell_len)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 50)

    wb.save(str(path))
    return f"已创建 Excel 文件：{path}\n工作表：{sheet_name}\n数据：{len(data)} 行 x {len(data[0]) if data else 0} 列"


@tool(
    name="read_workbook",
    description="读取 Excel 文件的内容，返回制表符分隔的文本表格。",
    schema={
        "type": "object",
        "properties": {
            "path": {
                "type": "string",
                "description": "Excel 文件路径",
            },
            "sheet_name": {
                "type": "string",
                "description": "要读取的工作表名称，默认读第一个",
            },
            "max_rows": {
                "type": "integer",
                "description": "最多读取行数，默认 100",
                "default": 100,
            },
        },
        "required": ["path"],
    },
)
def read_workbook(args: dict, ctx: ToolContext) -> str:
    if not _HAS_OPENPYXL:
        return _EXCEL_ERROR

    path = ctx.resolve(args["path"])
    if not path.exists():
        return f"[错误] 文件不存在：{path}"

    sheet_name = args.get("sheet_name")
    max_rows = int(args.get("max_rows", 100))

    wb = openpyxl.load_workbook(str(path), data_only=True)
    ws = wb[sheet_name] if sheet_name else wb.active

    rows = []
    for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row_idx > max_rows:
            break
        cells = [str(c) if c is not None else "" for c in row]
        rows.append("\t".join(cells))

    if not rows:
        return f"工作表 '{ws.title}' 为空"

    header_info = f"工作表：{ws.title} | 行数：{min(ws.max_row, max_rows)}"
    if ws.max_row > max_rows:
        header_info += f"（共 {ws.max_row} 行，已截断）"
    return header_info + "\n" + "\n".join(rows)


@tool(
    name="edit_cell",
    description="修改 Excel 文件中指定单元格的值或公式。",
    schema={
        "type": "object",
        "properties": {
            "path": {
                "type": "string",
                "description": "Excel 文件路径",
            },
            "sheet_name": {
                "type": "string",
                "description": "工作表名称，默认第一个",
            },
            "cell": {
                "type": "string",
                "description": "单元格位置，如 A1、B3",
            },
            "value": {
                "description": "要写入的值（字符串、数字或公式）",
            },
            "is_formula": {
                "type": "boolean",
                "description": "是否为公式（如 SUM(A1:A10)），默认 false",
                "default": False,
            },
        },
        "required": ["path", "cell", "value"],
    },
)
def edit_cell(args: dict, ctx: ToolContext) -> str:
    if not _HAS_OPENPYXL:
        return _EXCEL_ERROR

    path = ctx.resolve(args["path"])
    if not path.exists():
        return f"[错误] 文件不存在：{path}"

    sheet_name = args.get("sheet_name")
    cell_ref = args["cell"].upper()
    value = args["value"]
    is_formula = args.get("is_formula", False)

    wb = openpyxl.load_workbook(str(path))
    ws = wb[sheet_name] if sheet_name else wb.active

    if is_formula and isinstance(value, str) and not value.startswith("="):
        value = "=" + value

    ws[cell_ref] = value
    wb.save(str(path))
    return f"已更新 {ws.title}!{cell_ref} = {value}"


@tool(
    name="format_cells",
    description="设置 Excel 单元格区域的格式（加粗、字体颜色、背景色、数字格式）。",
    schema={
        "type": "object",
        "properties": {
            "path": {
                "type": "string",
                "description": "Excel 文件路径",
            },
            "sheet_name": {
                "type": "string",
                "description": "工作表名称，默认第一个",
            },
            "range": {
                "type": "string",
                "description": "单元格区域，如 A1:C3",
            },
            "bold": {
                "type": "boolean",
                "description": "是否加粗",
            },
            "font_color": {
                "type": "string",
                "description": "字体颜色，十六进制如 FF0000",
            },
            "bg_color": {
                "type": "string",
                "description": "背景色，十六进制如 FFFF00",
            },
            "number_format": {
                "type": "string",
                "description": "数字格式，如 0.00、#,##0、yyyy-mm-dd",
            },
        },
        "required": ["path", "range"],
    },
)
def format_cells(args: dict, ctx: ToolContext) -> str:
    if not _HAS_OPENPYXL:
        return _EXCEL_ERROR

    path = ctx.resolve(args["path"])
    if not path.exists():
        return f"[错误] 文件不存在：{path}"

    sheet_name = args.get("sheet_name")
    cell_range = args["range"].upper()
    bold = args.get("bold")
    font_color = args.get("font_color")
    bg_color = args.get("bg_color")
    number_format = args.get("number_format")

    wb = openpyxl.load_workbook(str(path))
    ws = wb[sheet_name] if sheet_name else wb.active

    font_kwargs = {}
    if bold is not None:
        font_kwargs["bold"] = bold
    if font_color:
        font_kwargs["color"] = font_color
    font = Font(**font_kwargs) if font_kwargs else None
    fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid") if bg_color else None

    changes = 0
    target = ws[cell_range]
    # 单元格 vs 区域：单个单元格返回 Cell，区域返回 tuple of tuples
    if isinstance(target, tuple):
        cells = [cell for row in target for cell in row]
    else:
        cells = [target]
    for cell in cells:
        if font:
            cell.font = font
        if fill:
            cell.fill = fill
        if number_format:
            cell.number_format = number_format
        changes += 1

    wb.save(str(path))
    parts = [f"已格式化 {ws.title}!{cell_range}（{changes} 个单元格）"]
    if bold:
        parts.append("加粗")
    if font_color:
        parts.append(f"字体色 #{font_color}")
    if bg_color:
        parts.append(f"背景色 #{bg_color}")
    if number_format:
        parts.append(f"数字格式 {number_format}")
    return "，".join(parts)


@tool(
    name="manage_sheet",
    description="管理工作表：创建、删除、重命名或列出所有工作表。",
    schema={
        "type": "object",
        "properties": {
            "path": {
                "type": "string",
                "description": "Excel 文件路径",
            },
            "action": {
                "type": "string",
                "description": "操作类型：create（创建）、delete（删除）、rename（重命名）、list（列出所有）",
                "enum": ["create", "delete", "rename", "list"],
            },
            "sheet_name": {
                "type": "string",
                "description": "目标工作表名称（create/delete/rename 时必填）",
            },
            "new_name": {
                "type": "string",
                "description": "新名称（rename 时必填）",
            },
        },
        "required": ["path", "action"],
    },
)
def manage_sheet(args: dict, ctx: ToolContext) -> str:
    if not _HAS_OPENPYXL:
        return _EXCEL_ERROR

    path = ctx.resolve(args["path"])
    if not path.exists():
        return f"[错误] 文件不存在：{path}"

    action = args["action"]
    sheet_name = args.get("sheet_name")
    new_name = args.get("new_name")

    wb = openpyxl.load_workbook(str(path))

    if action == "list":
        sheets = wb.sheetnames
        return f"工作表列表（{len(sheets)} 个）：\n" + "\n".join(f"  - {s}" for s in sheets)

    if not sheet_name:
        return "[错误] create/delete/rename 操作需要指定 sheet_name"

    if action == "create":
        if sheet_name in wb.sheetnames:
            return f"[错误] 工作表 '{sheet_name}' 已存在"
        wb.create_sheet(title=sheet_name)
        wb.save(str(path))
        return f"已创建工作表：{sheet_name}"

    if action == "delete":
        if sheet_name not in wb.sheetnames:
            return f"[错误] 工作表 '{sheet_name}' 不存在"
        if len(wb.sheetnames) <= 1:
            return "[错误] 不能删除唯一的工作表"
        del wb[sheet_name]
        wb.save(str(path))
        return f"已删除工作表：{sheet_name}"

    if action == "rename":
        if sheet_name not in wb.sheetnames:
            return f"[错误] 工作表 '{sheet_name}' 不存在"
        if not new_name:
            return "[错误] rename 操作需要指定 new_name"
        wb[sheet_name].title = new_name
        wb.save(str(path))
        return f"已将工作表 '{sheet_name}' 重命名为 '{new_name}'"

    return f"[错误] 未知操作：{action}，支持 create/delete/rename/list"


# ---------------------------------------------------------------------------
# 多 Agent 工具
# ---------------------------------------------------------------------------


@tool(
    name="spawn_agent",
    description="启动一个专业 Agent 来处理指定任务。Agent 在后台运行，立即返回 task 状态；只有主 Agent 可调用。",
    schema={
        "type": "object",
        "properties": {
            "agent": {
                "type": "string",
                "description": "Agent 名称，如 architect、coder、reviewer",
            },
            "task": {
                "type": "string",
                "description": "任务描述，越具体越好",
            },
            "context": {
                "type": "string",
                "description": "可选的额外上下文信息",
            },
            "depends_on": {
                "type": "array",
                "items": {"type": "string"},
                "description": "可选，上游 task_id 列表；未完成前本任务保持 pending",
            },
        },
        "required": ["agent", "task"],
    },
)
def spawn_agent(args: dict, ctx: ToolContext) -> str:
    if not ctx.can_spawn_agents:
        return json.dumps({
            "ok": False,
            "status": "failed",
            "error": "spawn_agent is only available to the main agent",
        }, ensure_ascii=False)

    if ctx.orchestrator is None:
        return json.dumps({
            "ok": False,
            "status": "failed",
            "error": "orchestrator is not configured",
        }, ensure_ascii=False)

    result = ctx.orchestrator.spawn_agent(
        agent_name=args["agent"],
        task_description=args["task"],
        context=args.get("context", ""),
        depends_on=args.get("depends_on", []),
    )
    return json.dumps(result, ensure_ascii=False)


@tool(
    name="check_agent_status",
    description="查询子 Agent 的执行状态。返回状态、结果或错误信息。",
    schema={
        "type": "object",
        "properties": {
            "task_id": {
                "type": "string",
                "description": "spawn_agent 返回的 task_id",
            },
        },
        "required": ["task_id"],
    },
)
def check_agent_status(args: dict, ctx: ToolContext) -> str:
    if ctx.orchestrator is None:
        return json.dumps({
            "ok": False,
            "status": "failed",
            "error": "orchestrator is not configured",
        }, ensure_ascii=False)

    result = ctx.orchestrator.check_status(args["task_id"])
    return json.dumps(result, ensure_ascii=False)


# ---------------------------------------------------------------------------
# 工具函数
# ---------------------------------------------------------------------------

def _human_size(size: int) -> str:
    for unit in ("B", "KB", "MB", "GB"):
        if size < 1024:
            return f"{size:.0f} {unit}"
        size //= 1024
    return f"{size:.0f} TB"
